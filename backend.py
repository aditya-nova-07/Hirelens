import os
import json
import time
from pathlib import Path

import streamlit as st
import httpx
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel
from pypdf import PdfReader
from docx import Document
from recommendations import ground_recommendations


# =========================
# 1. CONFIGURATION
# =========================

load_dotenv()

my_api_key = os.getenv("GROQ_API_KEY")
if not my_api_key:
    try:
        my_api_key = st.secrets.get("GROQ_API_KEY")
    except FileNotFoundError:
        # Local development can use .env without creating Streamlit secrets.
        my_api_key = None

if not my_api_key:
    st.error("GROQ_API_KEY is missing. Set it in .env or Streamlit app secrets.")
    st.stop()

# The inherited proxy settings can point to an unavailable local proxy. Direct
# HTTPS connectivity to Groq is available, so avoid inheriting those settings.
client = Groq(api_key=my_api_key, http_client=httpx.Client(trust_env=False))
model = "openai/gpt-oss-120b"


# =========================
# 2. PYDANTIC MODELS
# =========================

class JobD(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_requirements: list[str]
    responsibilities: list[str]


class MatchResult(BaseModel):
    score: float
    details: dict


class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []


class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None

    total_experience_years: float | None = None

    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []


jobd_schema = JobD.model_json_schema()
resume_schema = Resume.model_json_schema()


# =========================
# 3. JOB DESCRIPTION PARSER
# =========================

def parse_job_description(job_description):

    system_prompt = f"""
    You are an expert HR assistant.

    Extract structured information from the job description.

    Return ONLY valid JSON matching this schema:

    {jobd_schema}

    Do not return the schema itself.
    Do not invent information.

    If minimum experience is not mentioned, return null.
    If information for a list is missing, return an empty list.
    """

    messages = [
        {
            "role": "system",
            "content": system_prompt
        },
        {
            "role": "user",
            "content": f"Analyze this job description:\n{job_description}"
        }
    ]

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format={"type": "json_object"}
    )

    job_data = json.loads(
        response.choices[0].message.content
    )

    return JobD(**job_data)


# =========================
# 4. RESUME PARSER
# =========================

def parse_resume(resume_text):

    system_prompt = f"""
    You are an expert resume parser.

    Extract information from the resume based on its meaning,
    not only exact section headings.

    Different headings may contain relevant experience,
    including internships and work history.

    Return ONLY valid JSON matching this schema:

    {resume_schema}

    Rules:
    1. Do not invent information.
    2. If a value is unavailable, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    """

    messages = [
        {
            "role": "system",
            "content": system_prompt
        },
        {
            "role": "user",
            "content": f"Parse the following resume:\n{resume_text}"
        }
    ]

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format={"type": "json_object"}
    )

    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)

    return Resume(**data)


# =========================
# 5. RESUME-JOB MATCHING
# =========================

def final_score(job, resume):

    match_schema = MatchResult.model_json_schema()

    prompt = f"""
    You are an HR recruiter.

    Compare the candidate's resume with the job description.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}

    Return JSON matching this schema:

    {match_schema}

    Include these details inside the details field:
    1. Candidate name
    2. Matching skills
    3. Missing important skills
    4. Whether the experience requirement is met
    5. Overall match percentage from 0 to 100
    6. A short final verdict

    Keep the response concise and easy to read.
    """

    response = client.chat.completions.create(
        model=model,
        messages=[
            {
                "role": "user",
                "content": prompt
            }
        ],
        response_format={"type": "json_object"}
    )

    data = json.loads(
        response.choices[0].message.content
    )

    result = MatchResult(**data)

    # Keep the score within the expected range.
    if not 0 <= result.score <= 100:
        raise ValueError("The returned match score is outside 0–100.")

    return result


# =========================
# 6. PDF AND DOCX EXTRACTION
# =========================

def read_pdf(file_path):

    reader = PdfReader(file_path)
    text = ""

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


def read_docx(file_path):

    document = Document(file_path)
    text = ""

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"

    return text


def read_resume(file_path):

    # Streamlit UploadedFile exposes its filename as .name.
    suffix = Path(file_path.name).suffix.lower()

    file_path.seek(0)

    if suffix == ".pdf":
        return read_pdf(file_path)

    elif suffix == ".docx":
        return read_docx(file_path)

    return None


def classify_resume_document(resume_text):
    """Classify only documents that inexpensive resume checks could not resolve."""
    response = client.chat.completions.create(
        model=model,
        messages=[
            {
                "role": "system",
                "content": (
                    "Classify the supplied extracted document text as resume, non_resume, or uncertain. "
                    "A resume describes a person's work, education, skills, projects, or career profile. "
                    "Job advertisements, essays, school assignments, and unrelated reports are non_resume. "
                    "Do not infer missing content. Return only JSON with exactly these fields: "
                    '{"classification":"resume|non_resume|uncertain","reason":"short explanation"}.'
                ),
            },
            {"role": "user", "content": resume_text[:6000]},
        ],
        response_format={"type": "json_object"},
    )
    data = json.loads(response.choices[0].message.content)
    if (
        not isinstance(data, dict)
        or data.get("classification") not in {"resume", "non_resume", "uncertain"}
        or not isinstance(data.get("reason"), str)
        or not data["reason"].strip()
    ):
        raise ValueError("The document classifier returned an invalid response.")
    return {"classification": data["classification"], "reason": data["reason"].strip()[:300]}


def generate_skill_recommendations(resume_data, job_data, match_details):
    """Create and validate concise, job-grounded development recommendations."""
    allowed_resume_fields = {
        key: resume_data.get(key)
        for key in ("skills", "experiences", "education", "projects", "certifications")
        if isinstance(resume_data, dict) and key in resume_data
    }
    job_fields = {
        key: (job_data.get(key) if isinstance(job_data, dict) else getattr(job_data, key, None))
        for key in (
            "role", "required_skills", "preferred_skills", "minimum_experience",
            "education_requirements", "responsibilities",
        )
    }
    required = job_fields.get("required_skills") or []
    preferred = job_fields.get("preferred_skills") or []
    if not required and not preferred:
        return []

    prompt = f"""
Create concise skill-development recommendations for a candidate and this role.
Use only skills present in the supplied required_skills or preferred_skills lists.
Prioritize required skills. Do not claim a skill is absent from the candidate's abilities;
when it is not evidenced in resume data, say exactly that it is "Not demonstrated in the resume."
Only recommend further development for a demonstrated skill when there is a genuine evidence gap
and the category is "Area with insufficient evidence". Do not promise selection, invent credentials,
courses, resource links, or resume evidence. Return no more than 6 recommendations.
Estimated time must be a realistic estimate only when the supplied evidence supports one; otherwise null.

Return only JSON using this shape:
{{"recommendations":[{{"skill":"exact job skill","category":"Missing required skill|Missing preferred skill|Area with insufficient evidence","role_relevance":"why it matters based on the job data","candidate_evidence":"a short exact phrase from resume data, or Not demonstrated in the resume.","learning_action":"one concrete practice action or small project","priority":"High|Medium|Low","estimated_time":"estimate or null"}}]}}

JOB DATA:
{json.dumps(job_fields, ensure_ascii=False)}

CANDIDATE RESUME DATA:
{json.dumps(allowed_resume_fields, ensure_ascii=False)}

EXISTING MATCH DETAILS:
{json.dumps(match_details if isinstance(match_details, dict) else {}, ensure_ascii=False)}
"""
    response = client.chat.completions.create(
        model=model,
        messages=[{"role": "user", "content": prompt}],
        response_format={"type": "json_object"},
    )
    data = json.loads(response.choices[0].message.content)
    return ground_recommendations(data, allowed_resume_fields, job_fields, match_details)
